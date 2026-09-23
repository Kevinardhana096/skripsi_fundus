from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


source = Path(r"G:\My Drive\Skripsi\Outline_Penelitian_Sistem_Informasi_Cerdas_Fundus_revised_v3.docx")
output = Path(r"G:\My Drive\Skripsi\Outline_Penelitian_Sistem_Informasi_Cerdas_Fundus_revised_v4.docx")
copy2(source, output)
doc = Document(output)


def set_text(paragraph, value):
    paragraph.clear()
    paragraph.add_run(value)


def insert_after(paragraph, text):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.style = paragraph.style
    inserted.add_run(text)
    return inserted


replacements = {
    "Evaluasi Binary Cross-Entropy dan Asymmetric Loss untuk Klasifikasi Multi-Label Penyakit Mata pada Citra Fundus ODIR-5K dengan Analisis Grad-CAM":
        "Perbandingan Fungsi Loss pada Klasifikasi Multi-Label Penyakit Mata Menggunakan Citra Fundus ODIR-5K dengan Analisis Layer-CAM",
    "Research gap, novelty, backbone model, dan sebagian eksperimen belum dikunci; keputusan final menunggu literature review.":
        "Research gap, fungsi loss pembanding, dan metode XAI telah diperbarui; backbone model serta parameter pelatihan final ditetapkan setelah baseline awal.",
    "Arah penelitian: mengevaluasi Asymmetric Loss dan Explainable AI dalam peningkatan kualitas prediksi multi-label.":
        "Arah penelitian: membandingkan BCE, Weighted BCE, Focal Loss, dan Asymmetric Loss secara terkontrol, disertai interpretasi Layer-CAM.",
    "Urgensi data dan teknis: citra fundus dapat memuat lebih dari satu penyakit, sedangkan distribusi label pada ODIR-5K tidak seimbang. Model yang dilatih dengan Binary Cross-Entropy dapat lebih banyak menerima sinyal dari label negatif yang dominan, sehingga penyakit yang jarang muncul berisiko kurang terdeteksi. Karena itu, pengaruh Asymmetric Loss perlu diuji secara langsung dan terkontrol, bukan diasumsikan lebih baik.":
        "Urgensi data dan teknis: citra fundus dapat memuat lebih dari satu penyakit, sedangkan distribusi label pada ODIR-5K tidak seimbang. Model dapat lebih banyak menerima sinyal dari label negatif atau kelas mayoritas, sehingga penyakit yang jarang muncul berisiko kurang terdeteksi. Karena itu, BCE, Weighted BCE, Focal Loss, dan Asymmetric Loss perlu diuji secara langsung dan terkontrol, bukan diasumsikan memiliki kinerja terbaik.",
    "Urgensi kepercayaan: keluaran sistem kesehatan tidak cukup hanya berupa probabilitas. Pengguna perlu melihat bagian citra yang berkontribusi terhadap prediksi agar hasil dapat ditinjau secara kritis. Analisis Grad-CAM dapat membantu menunjukkan dasar visual prediksi, dengan posisi sebagai alat interpretasi pendukung dan bukan pengganti validasi klinis.":
        "Urgensi kepercayaan: keluaran sistem kesehatan tidak cukup hanya berupa probabilitas. Pengguna perlu melihat bagian citra yang berkontribusi terhadap prediksi agar hasil dapat ditinjau secara kritis. Layer-CAM digunakan untuk menunjukkan dasar visual prediksi dengan detail spasial yang lebih rinci melalui informasi dari beberapa layer CNN. Peta ini berfungsi sebagai interpretasi pendukung, bukan pengganti validasi klinis atau penanda lokasi lesi dari dokter.",
    "Urgensi akademik dan sistem informasi: penelitian terdahulu banyak menggabungkan arsitektur kompleks, resampling, attention, Transformer, atau knowledge distillation. Masih diperlukan evaluasi yang lebih terukur mengenai kontribusi loss function dengan backbone, pembagian data, dan konfigurasi pelatihan yang sama. Hasil ini diharapkan memberi dasar bagi pengembangan sistem pengolahan informasi visual yang lebih adil terhadap kelas minoritas dan lebih mudah ditinjau.":
        "Urgensi akademik dan sistem informasi: penelitian terdahulu banyak menggabungkan arsitektur kompleks, resampling, attention, Transformer, atau knowledge distillation. Masih diperlukan evaluasi yang lebih terukur mengenai kontribusi beberapa fungsi loss dengan backbone, pembagian data, dan konfigurasi pelatihan yang sama. Hasil ini diharapkan memberi dasar bagi pengembangan sistem pengolahan informasi visual yang lebih adil terhadap kelas minoritas dan lebih mudah ditinjau.",
    "Bagaimana pengaruh Asymmetric Loss terhadap kinerja model dibandingkan loss baseline?":
        "Bagaimana perbedaan kinerja BCE, Weighted BCE, Focal Loss, dan Asymmetric Loss pada klasifikasi multi-label?",
    "Bagaimana pengaruh Asymmetric Loss terhadap performa kelas mayoritas dan minoritas?":
        "Fungsi loss mana yang paling baik menjaga performa kelas mayoritas dan minoritas?",
    "Bagaimana Explainable AI dapat digunakan untuk menganalisis dasar visual dari prediksi model?":
        "Bagaimana Layer-CAM dapat digunakan untuk menganalisis dasar visual dari prediksi model?",
    "Mengevaluasi pengaruh Asymmetric Loss terhadap kinerja klasifikasi.":
        "Membandingkan kinerja BCE, Weighted BCE, Focal Loss, dan Asymmetric Loss secara terkontrol.",
    "Menganalisis performa model pada kelas mayoritas dan minoritas.":
        "Menganalisis performa setiap fungsi loss pada kelas mayoritas dan minoritas.",
    "Menghasilkan dan menganalisis penjelasan visual terhadap prediksi model menggunakan Explainable AI.":
        "Menghasilkan dan menganalisis penjelasan visual terhadap prediksi model menggunakan Layer-CAM.",
    "Perbandingan utama: BCE sebagai baseline dan Asymmetric Loss; loss tambahan mengikuti hasil literature review.":
        "Perbandingan utama: BCE, Weighted BCE, Focal Loss, dan Asymmetric Loss dengan konfigurasi pelatihan yang sama.",
    "XAI: Grad-CAM atau metode interpretabilitas yang dipilih berdasarkan kesesuaian dengan arsitektur.":
        "XAI: Layer-CAM pada layer konvolusi yang ditetapkan sebelum eksperimen; Grad-CAM dibahas sebagai pembanding konseptual di tinjauan pustaka.",
    "Pembagian data diupayakan pada level pasien untuk mengurangi potensi leakage antara mata kiri dan kanan pasien yang sama.":
        "Pembagian data dilakukan pada level pasien. Citra mata kiri dan kanan pasien yang sama selalu berada pada subset yang sama untuk mencegah data leakage.",
    "Grad-CAM / metode yang dipilih":
        "Layer-CAM: konsep, mekanisme, dan keterbatasan",
    "ODIR-5K: jumlah data, jenis citra, label, distribusi kelas":
        "ODIR-5K: data latih berlabel terdiri dari 3.500 pasien, citra fundus mata kiri dan kanan, serta delapan label penyakit/normal",
    "Fundus image → backbone CNN/arsitektur terpilih → feature representation → fully connected/multi-label head → sigmoid":
        "Citra mata kiri dan kanan → shared backbone CNN/arsitektur terpilih → feature representation → feature fusion → multi-label head → sigmoid",
    "Baseline: BCE":
        "BCE sebagai baseline standar",
    "Weighted BCE bila relevan":
        "Weighted BCE sebagai baseline pembobotan kelas",
    "Focal Loss bila relevan":
        "Focal Loss sebagai pembanding yang menekankan contoh sulit",
    "Asymmetric Loss sebagai metode utama":
        "Asymmetric Loss sebagai pembanding untuk positive-negative imbalance",
    "Prediksi → explanation map → perbandingan BCE vs ASL → analisis area perhatian model":
        "Prediksi per label → Layer-CAM per mata → analisis area perhatian model pada kasus representatif",
    "Apakah ASL meningkatkan aggregate metrics?":
        "Fungsi loss mana yang memberikan Macro-F1, Micro-F1, dan AUROC terbaik?",
    "Apakah kelas minoritas membaik?":
        "Apakah loss tertentu meningkatkan recall dan F1-score kelas minoritas?",
    "Bagaimana karakteristik attention/explanation map?":
        "Bagaimana karakteristik peta Layer-CAM pada prediksi benar dan salah?",
    "Hasil Explainable AI":
        "Hasil Layer-CAM",
    "Perbandingan visual BCE vs ASL":
        "Analisis visual Layer-CAM pada fungsi loss terpilih",
    "Literature Review → Identifikasi Research Gap → Penentuan Novelty → Pengumpulan Dataset → Data Understanding → Preprocessing → Patient-Level Split → Baseline Model → Eksperimen Loss Function → Evaluasi → Explainable AI → Analisis Hasil → Kesimpulan":
        "Literature Review → Identifikasi Research Gap → Penentuan Novelty → Pengumpulan Dataset → Data Understanding → Preprocessing → Patient-Level Split → Baseline Model → Eksperimen Fungsi Loss → Evaluasi → Layer-CAM → Analisis Hasil → Kesimpulan",
    "Metode XAI final":
        "Metode XAI final: Layer-CAM",
    "Asymmetric Loss dan Explainable AI saat ini merupakan kandidat kontribusi; novelty final mengikuti hasil review literatur.":
        "Perbandingan fungsi loss dan Layer-CAM merupakan rancangan kontribusi penelitian; kesimpulan tentang loss terbaik ditetapkan berdasarkan hasil eksperimen.",
}

for paragraph in doc.paragraphs:
    if paragraph.text in replacements:
        set_text(paragraph, replacements[paragraph.text])

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                if paragraph.text in replacements:
                    set_text(paragraph, replacements[paragraph.text])

for paragraph in list(doc.paragraphs):
    if paragraph.text == "Layer-CAM: konsep, mekanisme, dan keterbatasan":
        insert_after(
            paragraph,
            "Landasan pemilihan Layer-CAM: Jiang et al. (2021) menunjukkan bahwa Layer-CAM dapat memanfaatkan peta dari beberapa layer CNN untuk menghasilkan lokalisasi yang lebih rinci. Metode ini dipilih karena Grad-CAM pada layer akhir menghasilkan peta yang lebih kasar. Layer-CAM tetap diposisikan sebagai alat interpretasi kualitatif karena ODIR-5K tidak menyediakan anotasi lokasi lesi untuk memverifikasi heatmap secara klinis."
        )
        break

doc.save(output)
print(output)
